from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

REF = Path(r"C:\Users\sinny\OneDrive\文件\GitHub\spedu-labs\eng-grammar\tmp\numbers-docx\reference.docx")
OUT = Path(r"C:\Users\sinny\OneDrive\文件\GitHub\spedu-labs\eng-grammar\tmp\numbers-docx\topic03.數字1-100(待完成).docx")

EA_FONT = "台灣圓體-Light"
LATIN_FONT = "Comic Sans MS"
INK = "17233C"
BLUE = "45658A"
LIGHT_BLUE = "E8EEF6"
PALE = "F7F9FC"
GRID = "B9C7D8"


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_run(run, size=14, bold=False, color=INK):
    run.font.name = LATIN_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)
    rfonts.set(qn("w:eastAsia"), EA_FONT)


def fmt_paragraph(p, before=0, after=0, line=1.5, keep=False):
    f = p.paragraph_format
    f.space_before = Pt(before)
    f.space_after = Pt(after)
    f.line_spacing = line
    f.keep_with_next = keep


def add_text(doc, text="", size=14, bold=False, align=None, before=0, after=0, keep=False, color=INK):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    fmt_paragraph(p, before, after, 1.5, keep)
    set_run(p.add_run(text), size, bold, color)
    return p


def add_chapter(doc, title, first=False):
    if not first:
        doc.add_page_break()
    p = add_text(doc, title, 22, True, WD_ALIGN_PARAGRAPH.LEFT, after=3, keep=True)
    p.paragraph_format.keep_with_next = True
    p2 = add_text(doc, "班級：__________　姓名：______________", 14, False,
                  WD_ALIGN_PARAGRAPH.RIGHT, after=6, keep=True, color=BLUE)
    set_bottom_border(p2, color=BLUE, size="18")


def add_heading(doc, text):
    p = add_text(doc, text, 16, True, before=6, after=2, keep=True)
    set_bottom_border(p, color="9FB2C9", size="6")
    return p


def add_instruction(doc, text):
    return add_text(doc, text, 14, False, after=2, keep=True, color="3B4F68")


def add_q(doc, n, text, extra_line=False):
    p = add_text(doc, f"（{n}）{text}", 14, False, before=2, after=1)
    p.paragraph_format.line_spacing = 1.3
    if extra_line:
        add_answer_line(doc)
    return p


def add_answer_line(doc, width="80%"):
    p = add_text(doc, "　", 14, after=3)
    set_bottom_border(p, color="6D7F96", size="10")
    if width != "100%":
        p.paragraph_format.right_indent = Inches(1.0)
    return p


def set_bottom_border(p, color=GRID, size="6"):
    ppr = p._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRID, size="6"):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_cell_text(cell, text, size=14, bold=False, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    fmt_paragraph(p, 1, 1, 1.2, False)
    p.paragraph_format.left_indent = Inches(0)
    p.paragraph_format.right_indent = Inches(0)
    p.paragraph_format.first_line_indent = Inches(0)
    set_run(p.add_run(str(text)), size, bold, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def make_table(doc, rows, headers=None, widths=None, font_size=14):
    total = len(rows) + (1 if headers else 0)
    cols = len(headers) if headers else len(rows[0])
    table = doc.add_table(rows=total, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    idx = 0
    if headers:
        for c, value in enumerate(headers):
            set_cell_text(table.cell(0, c), value, font_size, True)
            shade_cell(table.cell(0, c), LIGHT_BLUE)
        trpr = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        trpr.append(repeat)
        idx = 1
    for r, values in enumerate(rows, start=idx):
        for c, value in enumerate(values):
            set_cell_text(table.cell(r, c), value, font_size)
            if (r - idx) % 2 == 1:
                shade_cell(table.cell(r, c), PALE)
    for row in table.rows:
        for c, cell in enumerate(row.cells):
            set_cell_border(cell)
            if widths:
                cell.width = Inches(widths[c])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_write_three(doc, numbers):
    rows = [[f"（{i}）{n}", "____________", "____________", "____________"]
            for i, n in enumerate(numbers, 1)]
    make_table(doc, rows, headers=["數字", "第 1 次", "第 2 次", "第 3 次"],
               widths=[1.15, 1.85, 1.85, 1.85])


def add_bingo(doc):
    table = doc.add_table(rows=5, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        row.height = Inches(0.52)
        for cell in row.cells:
            set_cell_text(cell, "", 16)
            set_cell_border(cell, color="7890AA", size="10")
            cell.width = Inches(1.35)
    doc.add_paragraph()


def build():
    doc = Document(REF)
    clear_body(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.59)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.59)
    section.right_margin = Inches(0.59)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.25)

    # 1
    add_chapter(doc, "1　我的數字提示板 Numbers 1–100", first=True)
    add_instruction(doc, "遇到不確定的數字時，先回到這一頁查一查。")
    rows = [
        ["1", "11", "10", "23"], ["2", "12", "20", "35"],
        ["3", "13", "30", "48"], ["4", "14", "40", "52"],
        ["5", "15", "50", "67"], ["6", "16", "60", "74"],
        ["7", "17", "70", "86"], ["8", "18", "80", "91"],
        ["9", "19", "90", "－"], ["10", "20", "100", "－"],
    ]
    make_table(doc, rows, ["1～10", "11～20", "整十", "組合數字"], [1.7, 1.7, 1.7, 1.7], 16)

    # 2
    add_chapter(doc, "2　Numbers 1–10")
    add_heading(doc, "2-1　寫一寫")
    add_instruction(doc, "照著數字寫三次。")
    add_write_three(doc, list(range(1, 11)))
    add_heading(doc, "2-2　聽力")
    add_instruction(doc, "聽老師念數字，寫下阿拉伯數字。")
    for i in range(1, 11): add_q(doc, i, "________________")
    add_heading(doc, "2-3　數字 ↔ 英文")
    for i, t in enumerate(["3 → __________________", "7 → __________________", "ten → ______", "one → ______", "5 → __________________", "nine → ______"], 1): add_q(doc, i, t)
    add_heading(doc, "2-4　圈一圈")
    for i, t in enumerate(["4　　four / five", "8　　six / eight", "two　　2 / 10", "seven　　7 / 9", "1　　one / three", "ten　　6 / 10"], 1): add_q(doc, i, t)
    add_heading(doc, "2-5　算一算")
    for i, t in enumerate(["3 + 4 = ______", "9 − 5 = ______", "2 + 6 = ______", "10 − 7 = ______", "1 + 8 = ______", "6 − 2 = ______"], 1): add_q(doc, i, t)
    add_heading(doc, "2-6　用一用")
    for i, t in enumerate(["Amy has ______ books.（3）", "Tom has ______ pencils.（7）", "I have ______ erasers.（2）", "There are ______ cats.（5）", "I see ______ birds.（8）", "He has ______ rulers.（4）"], 1): add_q(doc, i, t)

    # 3
    add_chapter(doc, "3　Numbers 11–20")
    add_heading(doc, "3-1　寫一寫")
    add_write_three(doc, list(range(11, 21)))
    add_heading(doc, "3-2　聽力")
    add_instruction(doc, "聽老師念數字，寫下阿拉伯數字。")
    for i in range(1, 11): add_q(doc, i, "________________")
    add_heading(doc, "3-3　數字 ↔ 英文")
    for i, t in enumerate(["13 → __________________", "20 → __________________", "eleven → ______", "sixteen → ______", "18 → __________________", "fourteen → ______"], 1): add_q(doc, i, t)
    add_heading(doc, "3-4　圈一圈")
    for i, t in enumerate(["12　　twelve / twenty", "15　　fifteen / fifty", "nineteen　　19 / 90", "thirteen　　13 / 30", "18　　eighteen / eighty", "twenty　　12 / 20"], 1): add_q(doc, i, t)
    add_heading(doc, "3-5　算一算")
    for i, t in enumerate(["11 + 4 = ______", "20 − 3 = ______", "13 + 5 = ______", "18 − 6 = ______", "12 + 7 = ______", "16 − 5 = ______"], 1): add_q(doc, i, t)
    add_heading(doc, "3-6　用一用")
    for i, t in enumerate(["There are ______ students.（12）", "Amy has ______ stickers.（15）", "We need ______ chairs.（20）", "Tom has ______ cards.（18）", "There are ______ books.（14）", "I see ______ stars.（19）"], 1): add_q(doc, i, t)

    # 4
    add_chapter(doc, "4　整十 Numbers 10–100")
    add_heading(doc, "4-1　看一看")
    add_instruction(doc, "觀察整十數字的規律，再完成後面的練習。")
    make_table(doc, [[10, "ten"], [20, "twenty"], [30, "thirty"], [40, "forty"], [50, "fifty"], [60, "sixty"], [70, "seventy"], [80, "eighty"], [90, "ninety"], [100, "one hundred"]], ["數字", "英文"], [2.2, 4.6])
    add_heading(doc, "4-2　寫一寫")
    add_write_three(doc, [10,20,30,40,50,60,70,80,90,100])
    add_heading(doc, "4-3　聽力")
    for i in range(1, 11): add_q(doc, i, "________________")
    add_heading(doc, "4-4　圈一圈")
    for i, t in enumerate(["30　　thirteen / thirty", "50　　fifteen / fifty", "eighty　　18 / 80", "ninety　　19 / 90", "40　　fourteen / forty", "seventy　　17 / 70"], 1): add_q(doc, i, t)
    add_heading(doc, "4-5　算一算")
    for i, t in enumerate(["20 + 30 = ______", "90 − 40 = ______", "10 + 60 = ______", "100 − 20 = ______", "30 + 50 = ______", "70 − 10 = ______"], 1): add_q(doc, i, t)
    add_heading(doc, "4-6　用一用")
    for i, t in enumerate(["The bag is ______ dollars.（50）", "There are ______ students.（30）", "We need ______ chairs.（40）", "The book is ______ dollars.（80）", "There are ______ cookies.（60）", "The class has ______ points.（100）"], 1): add_q(doc, i, t)

    # 5
    add_chapter(doc, "5　Numbers 20–29 找規律")
    add_heading(doc, "5-1　看規律")
    add_instruction(doc, "twenty 後面加上 1～9，就能組成 21～29。")
    make_table(doc, [[20,"twenty"],[21,"twenty-one"],[22,"twenty-two"],[23,"twenty-three"],[24,"twenty-four"],[25,"twenty-five"],[26,"twenty-six"],[27,"twenty-seven"],[28,"twenty-eight"],[29,"twenty-nine"]], ["數字", "英文"], [2.2,4.6])
    add_heading(doc, "5-2　補一補")
    for i, t in enumerate(["20, 21, ______, 23, ______", "twenty-four → ______", "27 → __________________", "twenty-nine → ______", "25, 26, ______, 28, ______", "twenty-two → ______"], 1): add_q(doc, i, t)
    add_heading(doc, "5-3　聽力")
    for i in range(1, 11): add_q(doc, i, "________________")
    add_heading(doc, "5-4　算一算")
    for i, t in enumerate(["20 + 3 = ______", "29 − 4 = ______", "21 + 6 = ______", "28 − 5 = ______", "24 + 5 = ______", "27 − 7 = ______"], 1): add_q(doc, i, t)
    add_heading(doc, "5-5　用一用")
    for i, t in enumerate(["Amy has ______ stickers.（23）", "The book is ______ dollars.（29）", "There are ______ desks.（25）", "Tom has ______ cards.（22）", "We need ______ pencils.（28）", "I see ______ stars.（26）"], 1): add_q(doc, i, t)

    # 6
    add_chapter(doc, "6　Numbers 21–99 我會組數字了")
    add_heading(doc, "6-1　組合規則")
    add_instruction(doc, "先說整十，再說個位數。例如：48 = forty-eight。")
    make_table(doc, [[23,"twenty-three"],[35,"thirty-five"],[48,"forty-eight"],[52,"fifty-two"],[67,"sixty-seven"],[74,"seventy-four"],[86,"eighty-six"],[91,"ninety-one"]], ["數字", "英文"], [2.2,4.6])
    add_heading(doc, "6-2　數字 ↔ 英文")
    for i, t in enumerate(["34 → __________________", "sixty-two → ______", "79 → __________________", "forty-five → ______", "88 → __________________", "ninety-three → ______", "51 → __________________", "seventy-six → ______"], 1): add_q(doc, i, t)
    add_heading(doc, "6-3　聽力")
    for i in range(1, 11): add_q(doc, i, "________________")
    add_heading(doc, "6-4　算一算")
    for i, t in enumerate(["30 + 4 = ______", "60 + 7 = ______", "80 + 6 = ______", "90 + 1 = ______", "50 + 2 = ______", "70 + 4 = ______", "40 + 8 = ______", "20 + 3 = ______"], 1): add_q(doc, i, t)
    h65 = add_heading(doc, "6-5　生活中的數字")
    h65.paragraph_format.page_break_before = True
    make_table(doc, [["notebook", "$35"],["pencil case", "$48"],["storybook", "$67"],["school bag", "$91"]], ["Item", "Price"], [4.1,2.7])
    for i, t in enumerate(["The notebook is ______ dollars.", "The pencil case is ______ dollars.", "The storybook is ______ dollars.", "The school bag is ______ dollars."], 1): add_q(doc, i, t)

    # 7
    add_chapter(doc, "7　Numbers 1–100 綜合挑戰")
    add_heading(doc, "7-1　混合挑戰")
    for i, t in enumerate(["17 → __________________", "forty-two → ______", "90 − 30 = ______"], 1): add_q(doc, i, t)
    add_heading(doc, "7-2　聽力綜合")
    add_instruction(doc, "聽老師念數字，寫下你聽到的阿拉伯數字。")
    for i in range(1, 11): add_q(doc, i, "________________")
    add_heading(doc, "7-3　生活數學")
    add_text(doc, "情境一｜文具店", 15, True, before=3, keep=True, color=BLUE)
    make_table(doc, [["pencil", "$8"],["eraser", "$12"],["notebook", "$35"]], ["Item", "Price"], [4.1,2.7])
    for i, t in enumerate(["A pencil and an eraser cost ______ dollars.", "A notebook and a pencil cost ______ dollars."], 1): add_q(doc, i, t)
    add_text(doc, "情境二｜班級點數", 15, True, before=5, keep=True, color=BLUE)
    make_table(doc, [["Amy", "28 points"],["Tom", "46 points"],["Kevin", "75 points"]], ["Student", "Points"], [4.1,2.7])
    for i, t in enumerate(["Amy and Tom have ______ points in all.", "Kevin has ______ more points than Amy."], 3): add_q(doc, i, t)
    h74 = add_heading(doc, "7-4　Numbers Bingo")
    h74.paragraph_format.page_break_before = True
    add_instruction(doc, "從 1～100 選 25 個數字填入格子。聽到數字時，把它圈起來。")
    add_bingo(doc)

    # Normalize section properties and remove inherited content from header/footer.
    for sec in doc.sections:
        sec.page_width = Inches(8.5)
        sec.page_height = Inches(11)
        sec.top_margin = Inches(0.59)
        sec.bottom_margin = Inches(0.55)
        sec.left_margin = Inches(0.59)
        sec.right_margin = Inches(0.59)
        for part in (sec.header, sec.footer):
            for p in part.paragraphs:
                p.text = ""

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
